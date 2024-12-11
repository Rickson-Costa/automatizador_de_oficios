from flask import Flask, render_template, request, send_file
from docx import Document
import pandas as pd
import os, zipfile, locale

app = Flask(__name__)
UPLOAD_FOLDER = "uploads"
GENERATED_FOLDER = "generated"
MODEL_FOLDER = "modelo"

# Diretórios necessários
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(GENERATED_FOLDER, exist_ok=True)

@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        # Upload da planilha e processamento
        file = request.files["spreadsheet"]
        if file:
            filepath = os.path.join(UPLOAD_FOLDER, file.filename)
            file.save(filepath)

            # Carregar modelo
            model_path = os.path.join(MODEL_FOLDER, "modelo.docx")
            if not os.path.exists(model_path):
                return "Erro: Arquivo modelo.docx não encontrado."

            # Processar planilha e gerar ofícios
            generate_oficios(filepath, model_path)

            # Compactar arquivos gerados
            zip_path = os.path.join(GENERATED_FOLDER, "oficios.zip")
            with zipfile.ZipFile(zip_path, 'w') as zipf:
                for root, _, files in os.walk(GENERATED_FOLDER):
                    for file in files:
                        if file.endswith(".docx"):
                            zipf.write(os.path.join(root, file), arcname=file)

            return send_file(zip_path, as_attachment=True)

    return render_template("index.html")



def generate_oficios(spreadsheet_path, model_path):
    # Configurar o locale para o formato brasileiro
    locale.setlocale(locale.LC_ALL, 'pt_BR.UTF-8')

    # Carregar dados da planilha, indicando que os nomes das colunas estão na terceira linha
    df = pd.read_excel(spreadsheet_path, header=2)
    df.columns = df.columns.str.strip()  # Remove espaços extras dos nomes das colunas

    # Verificar se as colunas necessárias existem
    required_columns = ['N', 'MUNICÍPIO', 'VLR. TOTAL']
    missing_columns = [col for col in required_columns if col not in df.columns]
    if missing_columns:
        raise ValueError(f"As seguintes colunas estão faltando na planilha: {', '.join(missing_columns)}")

    # Filtrar as colunas necessárias
    df = df[required_columns]

    # Iterar sobre cada linha da planilha
    for index, row in df.iterrows():
        # Verificar se algum valor necessário está vazio ou inválido
        if pd.isnull(row['N']) or pd.isnull(row['MUNICÍPIO']) or pd.isnull(row['VLR. TOTAL']):
            print(f"Pulando linha {index + 1}: dados incompletos.")
            continue

        try:
            numero_oficio = int(row['N'])  # Tenta converter 'N°' para inteiro
            prefeito_municipio = str(row['MUNICÍPIO'])
        except ValueError:
            print(f"Pulando linha {index + 1}: valor inválido em 'N'.")
            continue

        # Formatar o valor como moeda brasileira
        try:
            valor_formatado = locale.currency(float(row['VLR. TOTAL']), grouping=True).replace('R$', '').strip()
        except ValueError:
            print(f"Pulando linha {index + 1}: valor inválido em 'VLR. TOTAL'.")
            continue

        # Criar o documento
        doc = Document(model_path)

        # Substituir placeholders no modelo
        for paragraph in doc.paragraphs:
            paragraph.text = paragraph.text.replace("{{numero_oficio}}", str(numero_oficio))
            paragraph.text = paragraph.text.replace("{{prefeito_municipio}}", str(row['MUNICÍPIO']))
            paragraph.text = paragraph.text.replace("{{valor}}", valor_formatado)

        # Salvar o documento gerado
        output_path = os.path.join(GENERATED_FOLDER, f"Oficio de {prefeito_municipio}.docx")
        doc.save(output_path)
        print(f"Documento salvo: {output_path}")


if __name__ == "__main__":
    app.run(debug=True)